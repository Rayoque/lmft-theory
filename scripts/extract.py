# -*- coding: utf-8 -*-
"""Extract the five .docx packets to auditable text + JSON.
Walks body XML in document order so TABLES stay interleaved with paragraphs.
Preserves bold spans -- the source author used bold as her own emphasis."""
import json, os, glob, io, sys
import docx
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn

SRC = os.path.join(os.path.dirname(__file__), "..", "source")
OUT = os.path.join(os.path.dirname(__file__), "..", "extracted")

DOCMAP = {
    "Copy of Copy of Systems Theories Review.docx": "Systems Theories Review",
    "Copy of Copy of Post Modern Therapies Review.docx": "Post Modern Therapies Review",
    "Copy of Copy of Psychodynamic.docx": "Psychodynamic",
    "Copy of Copy of Humanistic.docx": "Humanistic/Existential Review",
    "Copy of Copy of CBT interventions notes.docx": "CBT interventions notes",
}

def iter_block_items(parent):
    body = parent.element.body
    for child in body.iterchildren():
        if child.tag == qn('w:p'):
            yield Paragraph(child, parent)
        elif child.tag == qn('w:tbl'):
            yield Table(child, parent)

def para_record(p, idx):
    runs = []
    for r in p.runs:
        if not r.text:
            continue
        runs.append({"text": r.text, "bold": bool(r.bold), "italic": bool(r.italic),
                     "underline": bool(r.underline)})
    bold_spans = [r["text"] for r in runs if r["bold"] and r["text"].strip()]
    # indent level for list paragraphs
    lvl = None
    try:
        numPr = p._p.pPr.numPr
        if numPr is not None and numPr.ilvl is not None:
            lvl = int(numPr.ilvl.val)
    except Exception:
        pass
    return {"i": idx, "kind": "p", "style": p.style.name, "level": lvl,
            "text": p.text.strip(), "bold_spans": bold_spans, "runs": runs}

def main():
    os.makedirs(OUT, exist_ok=True)
    manifest = []
    for path in sorted(glob.glob(os.path.join(SRC, "*.docx"))):
        base = os.path.basename(path)
        label = DOCMAP.get(base, base)
        d = docx.Document(path)
        blocks, idx = [], 0
        for blk in iter_block_items(d):
            if isinstance(blk, Paragraph):
                rec = para_record(blk, idx)
                if rec["text"]:
                    blocks.append(rec); idx += 1
            else:
                rows = []
                for row in blk.rows:
                    cells = []
                    for c in row.cells:
                        ctext = "\n".join(x.text.strip() for x in c.paragraphs if x.text.strip())
                        cbold = []
                        for x in c.paragraphs:
                            for r in x.runs:
                                if r.bold and r.text.strip():
                                    cbold.append(r.text)
                        cells.append({"text": ctext, "bold_spans": cbold})
                    rows.append(cells)
                blocks.append({"i": idx, "kind": "table", "rows": rows}); idx += 1

        outjson = os.path.join(OUT, label.replace("/", "-") + ".json")
        with io.open(outjson, "w", encoding="utf-8") as f:
            json.dump({"sourceDoc": label, "file": base, "blocks": blocks}, f,
                      ensure_ascii=False, indent=1)

        outtxt = os.path.join(OUT, label.replace("/", "-") + ".txt")
        with io.open(outtxt, "w", encoding="utf-8") as f:
            f.write("### %s\n### file: %s\n\n" % (label, base))
            for b in blocks:
                if b["kind"] == "p":
                    mark = "**" if b["bold_spans"] else "  "
                    lv = "" if b["level"] is None else ("  " * (b["level"] + 1))
                    f.write("[%04d]%s %s%s\n" % (b["i"], mark, lv, b["text"]))
                else:
                    f.write("[%04d]== TABLE ==\n" % b["i"])
                    for row in b["rows"]:
                        f.write("        | " + " | ".join(
                            c["text"].replace("\n", " / ") for c in row) + "\n")
        ntab = sum(1 for b in blocks if b["kind"] == "table")
        npar = sum(1 for b in blocks if b["kind"] == "p")
        nbold = sum(len(b.get("bold_spans", [])) for b in blocks if b["kind"] == "p")
        manifest.append({"sourceDoc": label, "file": base, "paragraphs": npar,
                         "tables": ntab, "boldSpans": nbold})
        print(f"{label:34s} paras={npar:4d} tables={ntab:2d} boldspans={nbold:4d}")

    with io.open(os.path.join(OUT, "_manifest.json"), "w", encoding="utf-8") as f:
        json.dump(manifest, f, ensure_ascii=False, indent=1)

if __name__ == "__main__":
    main()
